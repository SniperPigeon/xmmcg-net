from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import Permission
from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError
from rest_framework.test import APIClient

from songs.competitions import activate_competition, archive_competition
from songs.draws import join_draw, run_draw
from songs.exports import export_peer_review_scores
from songs.majdata_service import sync_chart
from songs.models import Competition, CompetitionPhase, PeerReview, Song
from songs.reviews import allocate_reviews, submit_reviews
from .test_contest import phase, upload_files
from .test_peer_review_allocation import make_chart, make_competition, make_user


class CompetitionTests(TestCase):
    def setUp(self):
        self.old = make_competition()
        self.user = make_user("participant")
        self.other = make_user("other-participant")
        self.own = make_chart(self.user, "旧作品甲")
        self.chart = make_chart(self.other, "旧作品乙")
        self.stage = phase()
        self.old.max_favorites = 1
        self.old.save()
        allocate_reviews(commit=True)
        submit_reviews(
            self.user,
            [
                {
                    "chart_id": self.chart.pk,
                    "score": "42.75",
                    "favorite": True,
                    "comment": "旧评语",
                }
            ],
        )
        self.api = APIClient()
        self.api.force_authenticate(self.user)

    def next_competition(self):
        archive_competition(self.old.pk)
        new = Competition.objects.create(
            name="下一届", max_favorites=1, reviews_per_chart=1
        )
        return activate_competition(new.pk)

    def test_archive_preserves_data_and_closes_all_public_operations(self):
        original_export = export_peer_review_scores(competition=self.old)
        review = PeerReview.objects.get()
        archive_competition(self.old.pk)
        self.old.refresh_from_db()
        self.assertEqual(self.old.status, "archived")
        self.assertIsNotNone(self.old.archived_at)
        review.refresh_from_db()
        self.assertEqual(review.comment, "旧评语")
        self.assertEqual(self.old.charts.count(), 2)
        before = load_workbook(BytesIO(original_export))
        after = load_workbook(BytesIO(export_peer_review_scores(competition=self.old)))
        for name in before.sheetnames:
            self.assertEqual(list(before[name].values), list(after[name].values))
        status = self.api.get("/api/contest/status/").data
        self.assertIsNone(status["competition"])
        self.assertEqual(status["submissions"], 0)
        self.assertFalse(any(status["capabilities"].values()))
        self.assertEqual(self.api.get("/api/contest/phases/").data, [])
        self.assertEqual(self.api.get("/api/charts/").data["count"], 0)
        self.assertEqual(self.api.get("/api/charts/me/").data["count"], 0)
        self.assertEqual(self.api.get("/api/reviews/tasks/").data["tasks"], [])
        self.assertEqual(self.api.get("/api/reviews/available/").data, [])
        for url, data in [
            ("/api/charts/", upload_files()),
            ("/api/songs/", {}),
            ("/api/reviews/submit/", {"reviews": []}),
        ]:
            self.assertEqual(self.api.post(url, data).status_code, 403)
        self.assertEqual(
            self.api.post(f"/api/charts/{self.chart.pk}/sync/").status_code, 404
        )
        with self.assertRaises(ValidationError):
            allocate_reviews(commit=True, competition_id=self.old.pk)
        with self.assertRaises(ValidationError):
            run_draw(self.stage.pk)
        with self.assertRaises(ValidationError):
            activate_competition(self.old.pk)

    def test_only_one_active_and_activation_requires_explicit_action(self):
        draft = Competition.objects.create(name="筹备比赛")
        self.assertEqual(draft.status, "draft")
        with self.assertRaises(ValidationError):
            activate_competition(draft.pk)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Competition.objects.create(name="绕过服务", status="active")
        archive_competition(self.old.pk)
        activated = activate_competition(draft.pk)
        self.assertIsNotNone(activated.activated_at)
        self.assertEqual(Competition.current().pk, draft.pk)

    def test_new_edition_excludes_old_phases_reviews_and_favorite_quota(self):
        phase("submission", offset=-1)
        new = self.next_competition()
        own = make_chart(self.user, "新作品甲", new)
        target = make_chart(self.other, "新作品乙", new)
        status = self.api.get("/api/contest/status/").data
        self.assertEqual(status["competition"]["id"], new.pk)
        self.assertTrue(status["capabilities"]["chart_upload"])
        self.assertFalse(status["capabilities"]["songs_enabled"])
        self.assertEqual(status["submissions"], 2)
        self.assertEqual(self.api.get("/api/contest/phases/").data, [])
        self.assertEqual(self.api.get("/api/reviews/tasks/").data["favorites_used"], 0)
        result = allocate_reviews(commit=True)
        self.assertEqual(result["new_tasks"], 2)
        self.assertEqual(
            {entry["count"] for entry in result["tasks_per_reviewer"]}, {1}
        )
        self.assertEqual(
            {
                row["chart_id"]
                for row in self.api.get("/api/reviews/tasks/").data["tasks"]
            },
            {target.pk},
        )
        submit_reviews(
            self.user, [{"chart_id": target.pk, "score": "33.25", "favorite": True}]
        )
        self.assertEqual(PeerReview.objects.count(), 2)
        with self.assertRaises(ValidationError):
            submit_reviews(self.other, [{"chart_id": self.own.pk, "score": 20}])
        current_rows = list(
            load_workbook(BytesIO(export_peer_review_scores()))["评分"].values
        )[2:]
        self.assertEqual({row[0] for row in current_rows}, {own.pk, target.pk})
        self.assertEqual(self.api.get(f"/api/charts/{self.chart.pk}/").status_code, 404)
        self.assertEqual(
            self.api.get(f"/api/charts/{self.chart.pk}/bundle/").status_code, 404
        )
        self.assertEqual(
            self.api.post(
                "/api/charts/", {"submission_key": str(self.chart.submission_key)}
            ).status_code,
            400,
        )

    @override_settings(MAX_SONGS_PER_USER=1)
    def test_song_quota_and_draw_results_reset_per_edition(self):
        def upload_song():
            return self.api.post(
                "/api/songs/",
                {
                    "title": "曲目",
                    "audio_file": SimpleUploadedFile("track.mp3", b"audio"),
                },
                format="multipart",
            )

        self.assertEqual(upload_song().status_code, 201)
        run_draw(self.stage.pk)
        new = self.next_competition()
        new_stage = phase(competition=new)
        # 阶段标识按届唯一，可沿用模板名称。
        new_stage.phase_key = self.stage.phase_key
        new_stage.save()
        self.assertEqual(self.api.get("/api/songs/draw/me/").data, [])
        self.assertEqual(upload_song().status_code, 201)
        self.assertEqual(upload_song().status_code, 400)
        self.assertEqual(len(self.api.get("/api/songs/").data), 1)
        join_draw(self.other)
        draw = run_draw(new_stage.pk)
        self.assertEqual(draw.entries.count(), 2)
        self.assertTrue(
            all(entry.song.competition_id == new.pk for entry in draw.entries.all())
        )
        self.assertEqual(Song.objects.filter(competition=self.old).count(), 1)
        with self.assertRaises(ValidationError):
            run_draw(self.stage.pk)

    @override_settings(ENABLE_CHART_FORWARD_TO_MAJDATA=True)
    def test_archived_net_retries_are_noop_and_inflight_sync_blocks_archive(self):
        self.chart.net_sync_status = "syncing"
        self.chart.save()
        with self.assertRaises(ValidationError):
            archive_competition(self.old.pk)
        self.chart.net_sync_status = "failed"
        self.chart.save()
        archive_competition(self.old.pk)
        with patch("songs.majdata_service.upload_chart") as upload:
            sync_chart(self.chart)
        upload.assert_not_called()
        self.assertEqual(self.chart.net_sync_status, "failed")

    def test_archive_during_object_upload_rejects_and_discards_new_assets(self):
        before = set(default_storage.objects)
        original_save = default_storage._save

        def archive_then_save(name, content):
            if Competition.current():
                archive_competition(self.old.pk)
            return original_save(name, content)

        with (
            patch.object(default_storage, "_save", side_effect=archive_then_save),
            patch("songs.chart_views.sync_chart") as sync,
        ):
            response = self.api.post("/api/charts/", upload_files(), format="multipart")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(set(default_storage.objects), before)
        self.assertEqual(self.old.charts.count(), 2)
        sync.assert_not_called()


class CompetitionAdminTests(TestCase):
    def setUp(self):
        self.admin = make_user("admin")
        self.admin.is_staff = self.admin.is_superuser = True
        self.admin.save()
        self.client.force_login(self.admin)
        self.competition = Competition.objects.create(
            name="后台测试", reviews_per_chart=2
        )

    def url(self, action, competition=None):
        return (
            f"/admin/songs/competition/{(competition or self.competition).pk}/{action}/"
        )

    def test_admin_create_copy_activate_archive_and_export(self):
        source = phase(competition=self.competition)
        self.assertContains(self.client.get(self.url("change")), "启用比赛")
        self.assertEqual(self.client.get(self.url("activate")).status_code, 200)
        self.competition.refresh_from_db()
        self.assertEqual(self.competition.status, "draft")
        self.assertEqual(self.client.post(self.url("activate")).status_code, 302)
        self.assertEqual(Competition.current().pk, self.competition.pk)
        make_chart(self.admin, competition=self.competition)
        self.assertContains(self.client.get(self.url("change")), "归档比赛")
        self.assertEqual(self.client.get(self.url("archive")).status_code, 200)
        self.assertIsNotNone(Competition.current())
        self.assertEqual(self.client.post(self.url("archive")).status_code, 302)
        self.assertIsNone(Competition.current())
        readonly = self.client.get(self.url("change"))
        self.assertContains(readonly, "导出本届成绩")
        self.assertNotContains(readonly, 'value="保存"')
        self.assertEqual(self.client.get(self.url("export")).status_code, 200)
        self.assertEqual(
            self.client.post(self.url("change"), {"name": "不应改名"}).status_code, 403
        )
        self.assertEqual(self.client.post(self.url("allocate")).status_code, 403)
        response = self.client.post(
            "/admin/songs/competition/add/",
            {
                "name": "复制的新比赛",
                "copy_from": self.competition.pk,
                "include_phases": "on",
                "submissions_open": "on",
                "reviews_open": "on",
                "reviews_per_chart": 8,
                "max_score": "50",
                "max_favorites": 3,
                "_save": "保存",
            },
        )
        self.assertEqual(
            response.status_code, 302, getattr(response, "context_data", None)
        )
        new = Competition.objects.get(name="复制的新比赛")
        self.assertEqual(new.status, "draft")
        self.assertEqual(new.reviews_per_chart, 2)
        self.assertEqual(new.charts.count(), 0)
        self.assertEqual(new.songs.count(), 0)
        copied = new.phases.get()
        self.assertEqual(copied.phase_key, source.phase_key)
        self.assertEqual(copied.start_time, source.start_time)
        self.assertFalse(copied.is_active)
        self.assertFalse(hasattr(copied, "draw"))
        self.assertEqual(
            self.client.post(
                f"/admin/songs/competitionphase/{copied.pk}/delete/", {"post": "yes"}
            ).status_code,
            302,
        )
        self.assertFalse(new.phases.exists())
        self.assertEqual(
            self.client.post(
                f"/admin/songs/competitionphase/{source.pk}/delete/", {"post": "yes"}
            ).status_code,
            403,
        )
        self.client.post(
            "/admin/songs/competitionphase/",
            {
                "action": "delete_selected",
                "_selected_action": [source.pk],
                "post": "yes",
            },
        )
        self.assertTrue(CompetitionPhase.objects.filter(pk=source.pk).exists())
        self.assertEqual(self.client.get(self.url("change", new)).status_code, 200)
        self.assertContains(
            self.client.get(f"/admin/songs/competitionphase/{source.pk}/change/"),
            source.name,
        )
        self.assertEqual(
            self.client.post(
                f"/admin/songs/competitionphase/{source.pk}/change/",
                {"name": "篡改历史"},
            ).status_code,
            403,
        )
        chart = self.competition.charts.get()
        self.assertEqual(
            self.client.post(
                f"/admin/songs/chart/{chart.pk}/change/", {"status": "archived"}
            ).status_code,
            403,
        )
        self.assertEqual(
            CompetitionPhase.objects.filter(competition=self.competition).get().name,
            source.name,
        )

    def test_staff_view_permission_cannot_change_lifecycle(self):
        viewer = make_user("viewer")
        viewer.is_staff = True
        viewer.save()
        viewer.user_permissions.add(Permission.objects.get(codename="view_competition"))
        self.client.force_login(viewer)
        self.assertEqual(self.client.get(self.url("change")).status_code, 200)
        self.assertEqual(self.client.get(self.url("export")).status_code, 200)
        for action in ("activate", "archive", "allocate"):
            self.assertEqual(self.client.post(self.url(action)).status_code, 403)
        self.assertEqual(
            self.client.post(
                "/admin/songs/competition/add/", {"name": "越权"}
            ).status_code,
            403,
        )

    def test_stale_admin_allocation_cannot_target_new_competition(self):
        activate_competition(self.competition.pk)
        archive_competition(self.competition.pk)
        new = Competition.objects.create(name="新届")
        activate_competition(new.pk)
        make_chart(self.admin, competition=new)
        self.assertEqual(self.client.post(self.url("allocate")).status_code, 403)
        self.assertFalse(new.charts.get().review_allocations.exists())

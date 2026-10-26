from decimal import Decimal
from io import BytesIO
from django.test import TestCase
from openpyxl import load_workbook
from songs.models import PeerReview
from songs.exports import export_peer_review_scores, trimmed_mean
from .test_peer_review_allocation import make_user, make_chart


class ExportTests(TestCase):
    def test_decimal_scores_comments_favorites_and_formula_escaping(self):
        chart = make_chart(make_user("author"), "=unsafe")
        for index, score in enumerate(["41.25", "42.50", "43.75"]):
            PeerReview.objects.create(
                chart=chart,
                reviewer=make_user(f"r{index}"),
                score=score,
                comment="=unsafe",
                favorite=index == 1,
            )
        workbook = load_workbook(BytesIO(export_peer_review_scores()))
        row = list(workbook["评分"].values)[2]
        self.assertEqual(row[1], "'=unsafe")
        self.assertEqual(row[-2:], (42.5, 1))
        self.assertEqual(list(workbook["评分与评语"].values)[3][4], "'=unsafe")

    def test_trimmed_mean(self):
        self.assertEqual(trimmed_mean([1, 2, 50]), Decimal("2.00"))
        self.assertEqual(trimmed_mean([1, 2]), Decimal("1.50"))
        self.assertIsNone(trimmed_mean([]))

    def test_original_chinese_layout_order_and_preferred_name(self):
        author = make_user("export-author")
        author.profile.preferred_name = "原有谱师名义"
        author.profile.save()
        first = make_chart(author, "第一首")
        first.designer = "文件内不同名义"
        first.save()
        second = make_chart(author, "第二首")
        reviewer = make_user("reviewer")
        PeerReview.objects.create(
            chart=first, reviewer=reviewer, score="10.25", comment="评语"
        )
        book = load_workbook(BytesIO(export_peer_review_scores()))
        scores = book["评分"]
        comments = book["评分与评语"]
        self.assertEqual(scores.cell(1, 1).value, "互评评分表")
        self.assertEqual(
            [scores.cell(2, col).value for col in (1, 2, 3)],
            ["谱面ID", "歌曲名", "谱师"],
        )
        self.assertEqual(
            [scores.cell(row, 1).value for row in (3, 4)], [first.pk, second.pk]
        )
        self.assertEqual(scores.cell(3, 3).value, "原有谱师名义")
        self.assertEqual(comments.cell(1, 1).value, "互评评分与评语表")
        self.assertEqual(
            [comments.cell(3, col).value for col in (4, 5)], ["评分", "评语"]
        )
        self.assertEqual(scores.freeze_panes, "D3")
        self.assertEqual(comments.freeze_panes, "D4")
        self.assertEqual(scores.cell(1, 1).fill.fgColor.rgb, "004472C4")
        self.assertEqual(trimmed_mean([1, "1.125", 50]), Decimal("1.12"))
